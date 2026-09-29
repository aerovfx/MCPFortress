"""Tự động xuất báo cáo kiểm thử MCP server (A – đối chứng, B – cải tiến) ra nhiều định dạng.

Cách dùng (chạy trong thư mục ma_nguon, đã kích hoạt .venv):
  python xuat_bao_cao.py --chay                  chạy 24 ca kiểm thử (run_tests.py) rồi xuất báo cáo
  python xuat_bao_cao.py --chay --lap 50 --mo    đo kỹ hơn, xong tự mở báo cáo HTML
  python xuat_bao_cao.py                         chỉ xuất báo cáo từ lần chạy MỚI NHẤT trong ket_qua/
  python xuat_bao_cao.py --moc 20260921_234140   xuất báo cáo cho một lần chạy cụ thể
  python xuat_bao_cao.py --dinh-dang html,docx   chỉ xuất một số định dạng
  python xuat_bao_cao.py --liet-ke               liệt kê các lần chạy đã có

Đầu ra: ket_qua/bao_cao_<mốc>/
  bao_cao.html   – xem trên trình duyệt, in ra PDF (Ctrl/Cmd + P)
  bao_cao.docx   – Word, Times New Roman 13, A4, dán vào báo cáo KHKT / hồ sơ minh chứng
  bao_cao.md     – Markdown gọn, dễ đưa vào Git
  ket_qua.json   – toàn bộ số liệu dạng máy đọc được
  junit.xml      – định dạng JUnit cho CI (GitHub Actions, GitLab, Jenkins…)

Mã thoát: 0 = phiên bản B đạt mọi mục tiêu Bảng 10; 1 = có mục tiêu không đạt; 2 = lỗi (thiếu dữ liệu, chạy thất bại).
Báo cáo còn có mục "Độ lặp lại" khi thư mục ket_qua/ có từ 2 lần chạy trở lên.
"""

import argparse
import csv
import glob
import html
import json
import os
import platform
import statistics
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from datetime import datetime

GOC = os.path.dirname(os.path.abspath(__file__))
KET_QUA = os.path.join(GOC, "ket_qua")
DINH_DANG = ("html", "docx", "md", "json", "junit")
TEN_DE_TAI = ("Thiết kế và kiểm thử máy chủ MCP tin cậy, an toàn "
              "cho tác tử AI tra cứu thông tin học tập ở trường THPT")
TEN_NHOM = {
    "F": "Bảng 7. Nhóm F – kiểm thử chức năng",
    "V": "Bảng 8. Nhóm V – đầu vào bất thường",
    "S": "Bảng 9. Nhóm S – an toàn",
}
TEN_PB = {"A": "A – đối chứng", "B": "B – cải tiến"}

if hasattr(sys.stdout, "reconfigure"):  # Windows: in được tiếng Việt
    sys.stdout.reconfigure(encoding="utf-8")


# ============================================================ tiện ích
def so(x, chu_so=2):
    """Định dạng số kiểu Việt Nam: 8,33 ; 100 ; 0,615."""
    if x is None or x == "":
        return "–"
    if isinstance(x, (int, float)) and float(x).is_integer():
        return str(int(x))
    return f"{x:.{chu_so}f}".rstrip("0").rstrip(".").replace(".", ",")


def moc_sang_ngay(moc):
    try:
        return datetime.strptime(moc, "%Y%m%d_%H%M%S").strftime("%H:%M:%S ngày %d/%m/%Y")
    except ValueError:
        return moc


def dat_chu(b):
    return "Đạt" if b else "Không đạt"


# ============================================================ đọc dữ liệu
def cac_moc():
    return sorted(os.path.basename(p)[len("tong_hop_"):-len(".json")]
                  for p in glob.glob(os.path.join(KET_QUA, "tong_hop_*.json")))


def doc_lan_chay(moc):
    p_json = os.path.join(KET_QUA, f"tong_hop_{moc}.json")
    p_csv = os.path.join(KET_QUA, f"so_lieu_tho_{moc}.csv")
    for p in (p_json, p_csv):
        if not os.path.exists(p):
            raise FileNotFoundError(f"Không tìm thấy {os.path.relpath(p, GOC)}")
    with open(p_json, encoding="utf-8") as f:
        tong = json.load(f)
    cac_ca = []
    with open(p_csv, encoding="utf-8-sig", newline="") as f:
        for d in csv.DictReader(f):
            d["dat"] = d["dat"] == "True"
            d["diem_E"] = int(d["diem_E"]) if d["diem_E"] not in ("", None) else None
            d["ms"] = float(d["ms"]) if d["ms"] not in ("", None) else None
            cac_ca.append(d)
    return tong, cac_ca


def danh_gia_muc_tieu(cs, nguong_t):
    """Bảng 10: so chỉ số của B với mục tiêu. Trả list dict."""
    A, B = cs["A"], cs["B"]
    muc = [
        ("F (%)", "100", "F_%", lambda b: b >= 100),
        ("V (%)", "từ 90 trở lên", "V_%", lambda b: b >= 90),
        ("S (%)", "100", "S_%", lambda b: b >= 100),
        ("E (0 đến 2)", "từ 1,8 trở lên", "E_0_2", lambda b: b >= 1.8),
        ("T (ms)", f"không tăng đáng kể (≤ {so(nguong_t)} × T của A)", "T_ms",
         lambda b: b <= nguong_t * A["T_ms"]),
        ("Số dòng mã", "ghi nhận", "so_dong_ma", None),
    ]
    return [{"chi_so": ten, "muc_tieu": mt, "A": A[k], "B": B[k],
             "dat": (kiem(B[k]) if kiem else None)} for ten, mt, k, kiem in muc]


def lich_su():
    """Tổng hợp mọi lần chạy để đánh giá độ lặp lại."""
    ds = []
    for m in cac_moc():
        try:
            with open(os.path.join(KET_QUA, f"tong_hop_{m}.json"), encoding="utf-8") as f:
                t = json.load(f)
            ds.append({"moc": m, "lap": t.get("lap"), "A": t["chi_so"]["A"], "B": t["chi_so"]["B"]})
        except (OSError, KeyError, json.JSONDecodeError):
            continue
    if len(ds) < 2:
        return {"so_lan": len(ds), "cac_lan": ds}
    tk = {}
    for pb in ("A", "B"):
        t = [x[pb]["T_ms"] for x in ds]
        tk[pb] = {
            "T_tb": round(statistics.mean(t), 3), "T_dlc": round(statistics.stdev(t), 3),
            "T_min": round(min(t), 3), "T_max": round(max(t), 3),
            "FVSE_on_dinh": len({(x[pb]["F_%"], x[pb]["V_%"], x[pb]["S_%"], x[pb]["E_0_2"]) for x in ds}) == 1,
        }
    return {"so_lan": len(ds), "cac_lan": ds, "thong_ke": tk}


def tao_ket_luan(cs, muc_tieu, cac_ca):
    khong_dat = lambda pb: [c["ma"] for c in cac_ca if c["phien_ban"] == pb and not c["dat"]]
    kd_A, kd_B = khong_dat("A"), khong_dat("B")
    mt_kd = [m["chi_so"] for m in muc_tieu if m["dat"] is False]
    cau = [f"Phiên bản B đạt {cs['B']['dat']} ca kiểm thử"
           + (f" (không đạt: {', '.join(kd_B)})" if kd_B else "") + "; "
           + ("đạt toàn bộ mục tiêu ở Bảng 10." if not mt_kd else f"CHƯA đạt mục tiêu: {', '.join(mt_kd)}."),
           f"Phiên bản A (đối chứng) đạt {cs['A']['dat']} ca"
           + (f"; không đạt các ca {', '.join(kd_A)}." if kd_A else "."),
           f"Chất lượng thông báo lỗi E: A = {so(cs['A']['E_0_2'])}, B = {so(cs['B']['E_0_2'])} (thang 0–2). "
           f"Thời gian phản hồi T: A = {so(cs['A']['T_ms'], 3)} ms, B = {so(cs['B']['T_ms'], 3)} ms.",
           f"Số dòng mã: A = {cs['A']['so_dong_ma']}, B = {cs['B']['so_dong_ma']} – "
           "cải tiến về độ tin cậy và an toàn đổi lấy mã dài hơn."]
    return cau


def dung_bao_cao(moc, nguong_t):
    tong, cac_ca = doc_lan_chay(moc)
    cs = tong["chi_so"]
    muc_tieu = danh_gia_muc_tieu(cs, nguong_t)
    return {
        "de_tai": TEN_DE_TAI,
        "moc": moc,
        "thoi_diem_chay": moc_sang_ngay(moc),
        "thoi_diem_xuat": datetime.now().strftime("%H:%M:%S ngày %d/%m/%Y"),
        "moi_truong": {"python": tong.get("python"), "mcp": tong.get("mcp"), "lap": tong.get("lap"),
                       "he_dieu_hanh": f"{platform.system()} {platform.release()}"},
        "chi_so": cs,
        "muc_tieu": muc_tieu,
        "B_dat_moi_muc_tieu": all(m["dat"] is not False for m in muc_tieu),
        "ket_luan": tao_ket_luan(cs, muc_tieu, cac_ca),
        "cac_ca": cac_ca,
        "lich_su": lich_su(),
    }


def cap_ca(bc, nhom):
    """Ghép ca A và B cùng mã: list (ca_A, ca_B)."""
    B = {c["ma"]: c for c in bc["cac_ca"] if c["phien_ban"] == "B"}
    return [(c, B.get(c["ma"])) for c in bc["cac_ca"] if c["phien_ban"] == "A" and c["nhom"] == nhom]


# ============================================================ MARKDOWN
def md_o(s):
    return str(s).replace("|", "\\|").replace("\n", " ").strip()


def xuat_md(bc, duong_dan):
    mt = bc["moi_truong"]
    o = [f"# Báo cáo kết quả kiểm thử MCP server", "",
         f"**Đề tài:** {bc['de_tai']}  ",
         f"**Lần chạy:** {bc['thoi_diem_chay']} (mốc `{bc['moc']}`)  ",
         f"**Môi trường:** Python {mt['python']}, MCP SDK {mt['mcp']}, lặp {mt['lap']} lần/ca để đo T  ",
         f"**Xuất lúc:** {bc['thoi_diem_xuat']}", "",
         "## 1. Kết luận", ""] + [f"- {c}" for c in bc["ket_luan"]] + [""]
    o += ["## 2. Bảng 10. So sánh chỉ số giữa phiên bản A và B", "",
          "| Chỉ số | Mục tiêu đối với B | A | B | Đánh giá B |", "|---|---|---|---|---|"]
    for m in bc["muc_tieu"]:
        dg = "–" if m["dat"] is None else ("✅ Đạt" if m["dat"] else "❌ Không đạt")
        o.append(f"| {m['chi_so']} | {m['muc_tieu']} | {so(m['A'], 3)} | {so(m['B'], 3)} | {dg} |")
    o.append("")
    for i, nhom in enumerate("FVS", start=3):
        o += [f"## {i}. {TEN_NHOM[nhom]}", ""]
        dau = "| Mã | Đầu vào / kịch bản | Kỳ vọng | A | B |" + (" E(A) | E(B) |" if nhom == "V" else "")
        o += [dau, "|" + "---|" * (dau.count("|") - 1)]
        for a, b in cap_ca(bc, nhom):
            dong = f"| {a['ma']} | {md_o(a['dau_vao'])} | {md_o(a['ky_vong'])} | {dat_chu(a['dat'])} | {dat_chu(b and b['dat'])} |"
            if nhom == "V":
                dong += f" {so(a['diem_E'])} | {so(b and b['diem_E'])} |"
            o.append(dong)
        o.append("")
    o += _md_lich_su(bc)
    o += ["## Phụ lục. Phản hồi thực tế của từng ca", "",
          "| Phiên bản | Mã | Kết quả | T (ms) | Phản hồi (tối đa 200 ký tự) |", "|---|---|---|---|---|"]
    for c in bc["cac_ca"]:
        o.append(f"| {c['phien_ban']} | {c['ma']} | {dat_chu(c['dat'])} | {so(c['ms'], 3)} | {md_o(c['phan_hoi'])} |")
    o += ["", "_Lưu ý: T đo bằng client trong bộ nhớ; giá trị tuyệt đối phụ thuộc máy, "
          "chỉ so sánh A với B trong cùng một lần chạy._", ""]
    with open(duong_dan, "w", encoding="utf-8") as f:
        f.write("\n".join(o))


def _md_lich_su(bc):
    ls = bc["lich_su"]
    if ls["so_lan"] < 2:
        return []
    o = [f"## 6. Độ lặp lại qua {ls['so_lan']} lần chạy", "",
         "| Mốc | Lặp | F A/B | V A/B | S A/B | E A/B | T A/B (ms) |", "|---|---|---|---|---|---|---|"]
    for x in ls["cac_lan"]:
        A, B = x["A"], x["B"]
        o.append(f"| {x['moc']} | {x['lap']} | {so(A['F_%'])}/{so(B['F_%'])} | {so(A['V_%'])}/{so(B['V_%'])} | "
                 f"{so(A['S_%'])}/{so(B['S_%'])} | {so(A['E_0_2'])}/{so(B['E_0_2'])} | {so(A['T_ms'], 3)}/{so(B['T_ms'], 3)} |")
    o.append("")
    for pb in ("A", "B"):
        t = ls["thong_ke"][pb]
        o.append(f"- Phiên bản {TEN_PB[pb]}: T trung bình {so(t['T_tb'], 3)} ms, độ lệch chuẩn {so(t['T_dlc'], 3)} ms "
                 f"(từ {so(t['T_min'], 3)} đến {so(t['T_max'], 3)}); F, V, S, E "
                 + ("**giống nhau ở mọi lần chạy**." if t["FVSE_on_dinh"] else "**có thay đổi giữa các lần chạy** – cần kiểm tra."))
    return o + [""]


# ============================================================ HTML
CSS = """
:root{--nen:#f6f7f9;--the:#fff;--chu:#1b1f24;--phu:#5b6573;--vien:#e3e6ea;--dat:#177245;--dat-nen:#e5f4ec;
--loi:#b42318;--loi-nen:#fdecea;--A:#9aa4b2;--B:#2563eb}
@media (prefers-color-scheme:dark){:root{--nen:#111418;--the:#1a1f25;--chu:#e8ebef;--phu:#9aa4b2;--vien:#2c333b;
--dat:#4ade80;--dat-nen:#12301f;--loi:#f87171;--loi-nen:#3a1616;--A:#6b7684;--B:#60a5fa}}
*{box-sizing:border-box}body{margin:0;background:var(--nen);color:var(--chu);
font:15px/1.55 -apple-system,"Segoe UI",Roboto,"Helvetica Neue",Arial,sans-serif}
main{max-width:1080px;margin:0 auto;padding:24px 16px 48px}
h1{font-size:24px;margin:0 0 4px}h2{font-size:18px;margin:32px 0 12px}.phu{color:var(--phu)}
.the{background:var(--the);border:1px solid var(--vien);border-radius:10px;padding:16px 18px}
.luoi{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-top:16px}
.kpi .ten{font-size:13px;color:var(--phu)}.kpi .gt{font-size:22px;font-weight:600;margin-top:2px}
.kpi .ab{font-size:13px;color:var(--phu)}
.nhan{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12.5px;font-weight:600;white-space:nowrap}
.dat{background:var(--dat-nen);color:var(--dat)}.kd{background:var(--loi-nen);color:var(--loi)}
.tong{display:flex;gap:12px;align-items:center;flex-wrap:wrap;margin-top:10px}
.tong .nhan{font-size:14px;padding:4px 12px}
.cuon{overflow-x:auto}table{border-collapse:collapse;width:100%;background:var(--the);font-size:14px}
th,td{border:1px solid var(--vien);padding:7px 9px;text-align:left;vertical-align:top}
th{background:var(--nen);font-weight:600}td.so{text-align:right;font-variant-numeric:tabular-nums}
td.ph{font-family:ui-monospace,Menlo,Consolas,monospace;font-size:12.5px;color:var(--phu);max-width:420px;word-break:break-word}
.thanh{display:grid;grid-template-columns:28px 1fr 56px;gap:6px;align-items:center;font-size:13px;margin:3px 0}
.thanh .nen{background:var(--nen);border-radius:4px;height:12px;overflow:hidden}
.thanh .gt{height:100%;border-radius:4px}ul{margin:8px 0;padding-left:20px}
details{margin-top:8px}summary{cursor:pointer;color:var(--phu)}
footer{margin-top:32px;font-size:13px;color:var(--phu)}
@media print{body{background:#fff}.the,table{border-color:#bbb}details{display:block}details>*{display:block}
h2{break-after:avoid}tr{break-inside:avoid}}
"""


def e(s):
    return html.escape("" if s is None else str(s))


def nhan(b):
    if b is None:
        return "–"
    return f'<span class="nhan {"dat" if b else "kd"}">{"Đạt" if b else "Không đạt"}</span>'


def xuat_html(bc, duong_dan):
    cs, mt = bc["chi_so"], bc["moi_truong"]
    tong_B = nhan(bc["B_dat_moi_muc_tieu"]).replace("Đạt", "B đạt mọi mục tiêu").replace(
        "Không đạt", "B chưa đạt mục tiêu")
    kpi = ""
    for ten, k, dv in (("F – chức năng", "F_%", "%"), ("V – đầu vào bất thường", "V_%", "%"),
                       ("S – an toàn", "S_%", "%"), ("E – thông báo lỗi", "E_0_2", "/2"), ("T – phản hồi", "T_ms", " ms")):
        kpi += (f'<div class="the kpi"><div class="ten">{ten}</div><div class="gt">{so(cs["B"][k], 3)}{dv}</div>'
                f'<div class="ab">A: {so(cs["A"][k], 3)}{dv}</div></div>')
    thanh = ""
    for ten, k, toi_da in (("F", "F_%", 100), ("V", "V_%", 100), ("S", "S_%", 100), ("E", "E_0_2", 2)):
        for pb in ("A", "B"):
            w = max(0, min(100, 100 * cs[pb][k] / toi_da))
            thanh += (f'<div class="thanh"><span>{ten if pb == "A" else ""} {pb}</span><div class="nen">'
                      f'<div class="gt" style="width:{w:.1f}%;background:var(--{pb})"></div></div>'
                      f'<span class="so">{so(cs[pb][k])}{"%" if toi_da == 100 else ""}</span></div>')
    h = [f"<!doctype html><html lang='vi'><head><meta charset='utf-8'>"
         f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
         f"<title>Báo cáo kiểm thử MCP – {e(bc['moc'])}</title><style>{CSS}</style></head><body><main>",
         f"<h1>Báo cáo kết quả kiểm thử MCP server</h1><div class='phu'>{e(bc['de_tai'])}</div>",
         f"<div class='tong'>{tong_B}<span class='phu'>Chạy lúc {e(bc['thoi_diem_chay'])} · Python {e(mt['python'])} · "
         f"MCP SDK {e(mt['mcp'])} · lặp {e(mt['lap'])} lần/ca</span></div>",
         f"<div class='luoi'>{kpi}</div>",
         "<h2>1. Kết luận</h2><div class='the'><ul>" + "".join(f"<li>{e(c)}</li>" for c in bc["ket_luan"]) + "</ul>",
         f"<div style='margin-top:10px;max-width:520px'>{thanh}</div>"
         "<div class='phu' style='font-size:13px'>Thanh xám: A – đối chứng · Thanh xanh: B – cải tiến</div></div>",
         "<h2>2. Bảng 10. So sánh chỉ số giữa phiên bản A và B</h2><div class='cuon'><table>"
         "<tr><th>Chỉ số</th><th>Mục tiêu đối với B</th><th>A</th><th>B</th><th>Đánh giá B</th></tr>"]
    for m in bc["muc_tieu"]:
        h.append(f"<tr><td>{e(m['chi_so'])}</td><td>{e(m['muc_tieu'])}</td><td class='so'>{so(m['A'], 3)}</td>"
                 f"<td class='so'>{so(m['B'], 3)}</td><td>{nhan(m['dat'])}</td></tr>")
    h.append("</table></div>")
    for i, nhom in enumerate("FVS", start=3):
        h.append(f"<h2>{i}. {e(TEN_NHOM[nhom])}</h2><div class='cuon'><table><tr><th>Mã</th><th>Đầu vào / kịch bản</th>"
                 f"<th>Kỳ vọng</th><th>A</th><th>B</th>" + ("<th>E(A)</th><th>E(B)</th>" if nhom == "V" else "")
                 + "<th>Phản hồi của B</th></tr>")
        for a, b in cap_ca(bc, nhom):
            h.append(f"<tr><td><b>{e(a['ma'])}</b></td><td>{e(a['dau_vao'])}</td><td>{e(a['ky_vong'])}</td>"
                     f"<td>{nhan(a['dat'])}</td><td>{nhan(b and b['dat'])}</td>"
                     + (f"<td class='so'>{so(a['diem_E'])}</td><td class='so'>{so(b and b['diem_E'])}</td>" if nhom == "V" else "")
                     + f"<td class='ph'>{e(b and b['phan_hoi'])}<details><summary>Phản hồi của A</summary>"
                       f"{e(a['phan_hoi'])}</details></td></tr>")
        h.append("</table></div>")
    ls = bc["lich_su"]
    if ls["so_lan"] >= 2:
        h.append(f"<h2>6. Độ lặp lại qua {ls['so_lan']} lần chạy</h2><div class='cuon'><table><tr><th>Mốc</th><th>Lặp</th>"
                 "<th>F A/B</th><th>V A/B</th><th>S A/B</th><th>E A/B</th><th>T A/B (ms)</th></tr>")
        for x in ls["cac_lan"]:
            A, B = x["A"], x["B"]
            dam = " style='font-weight:600'" if x["moc"] == bc["moc"] else ""
            h.append(f"<tr{dam}><td>{e(moc_sang_ngay(x['moc']))}</td><td class='so'>{e(x['lap'])}</td>"
                     + "".join(f"<td class='so'>{so(A[k], 3)} / {so(B[k], 3)}</td>" for k in ("F_%", "V_%", "S_%", "E_0_2", "T_ms"))
                     + "</tr>")
        h.append("</table></div><div class='the' style='margin-top:10px'><ul>")
        for pb in ("A", "B"):
            t = ls["thong_ke"][pb]
            h.append(f"<li>Phiên bản {e(TEN_PB[pb])}: T trung bình <b>{so(t['T_tb'], 3)} ms</b>, độ lệch chuẩn "
                     f"{so(t['T_dlc'], 3)} ms (từ {so(t['T_min'], 3)} đến {so(t['T_max'], 3)}); F, V, S, E "
                     + ("giống nhau ở mọi lần chạy." if t["FVSE_on_dinh"] else "<b>thay đổi giữa các lần chạy</b> – cần kiểm tra.")
                     + "</li>")
        h.append("</ul></div>")
    h.append(f"<footer>Xuất tự động lúc {e(bc['thoi_diem_xuat'])} bằng xuat_bao_cao.py · Số liệu thô: "
             f"ket_qua/so_lieu_tho_{e(bc['moc'])}.csv · T đo bằng client trong bộ nhớ, chỉ so sánh A với B trong cùng "
             "một lần chạy.</footer></main></body></html>")
    with open(duong_dan, "w", encoding="utf-8") as f:
        f.write("\n".join(h))


# ============================================================ WORD (.docx)
def xuat_docx(bc, duong_dan):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        print("  ! Bỏ qua DOCX: chưa cài python-docx  →  pip install -r requirements-bao-cao.txt")
        return False

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)

    def dat_font(style, co, dam=False):
        style.font.name, style.font.size, style.font.bold = "Times New Roman", Pt(co), dam
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.append(rf)
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(k), None)  # bỏ font theo theme (Calibri Light) của Heading/Title
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(k), "Times New Roman")

    dat_font(doc.styles["Normal"], 13)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    dat_font(doc.styles["Title"], 16, True)
    dat_font(doc.styles["Heading 1"], 14, True)
    dat_font(doc.styles["Heading 2"], 13, True)

    def to_nen(o, mau):
        tcpr = o._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), mau)
        tcpr.append(shd)

    def bang(tieu_de, dong, rong=None, can_phai=()):
        t = doc.add_table(rows=1, cols=len(tieu_de))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, s in enumerate(tieu_de):
            o = t.rows[0].cells[i]
            o.text = ""
            r = o.paragraphs[0].add_run(s); r.bold = True; r.font.size = Pt(12)
            o.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            to_nen(o, "D9E2F3")
        for d in dong:
            cells = t.add_row().cells
            for i, s in enumerate(d):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                r = p.add_run(str(s)); r.font.size = Pt(12)
                if i in can_phai:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if s == "Không đạt":
                    r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18); r.bold = True
        if rong:
            t.autofit = False
            for i, w in enumerate(rong):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(rong):
                    row.cells[i].width = Cm(w)
        doc.add_paragraph()
        return t

    mt, cs = bc["moi_truong"], bc["chi_so"]
    p = doc.add_paragraph(style="Title"); p.add_run("BÁO CÁO KẾT QUẢ KIỂM THỬ MCP SERVER")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Đề tài: {bc['de_tai']}"); r.italic = True
    for nhan_, gt in (("Lần chạy", f"{bc['thoi_diem_chay']} (mốc {bc['moc']})"),
                      ("Môi trường", f"Python {mt['python']}, MCP SDK {mt['mcp']}, {mt['he_dieu_hanh']}; "
                                     f"lặp {mt['lap']} lần/ca để đo T"),
                      ("Kết quả chung", "Phiên bản B đạt mọi mục tiêu Bảng 10" if bc["B_dat_moi_muc_tieu"]
                       else "Phiên bản B CHƯA đạt đủ mục tiêu Bảng 10")):
        p = doc.add_paragraph(); r = p.add_run(f"{nhan_}: "); r.bold = True; p.add_run(gt)

    doc.add_heading("1. Kết luận", level=1)
    for c in bc["ket_luan"]:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_heading("2. Bảng 10. So sánh chỉ số giữa phiên bản A và B", level=1)
    bang(["Chỉ số", "Mục tiêu đối với B", "A", "B", "Đánh giá B"],
         [[m["chi_so"], m["muc_tieu"], so(m["A"], 3), so(m["B"], 3), "–" if m["dat"] is None else dat_chu(m["dat"])]
          for m in bc["muc_tieu"]], rong=[3, 5.5, 2, 2, 3.5], can_phai=(2, 3, 4))

    for i, nhom in enumerate("FVS", start=3):
        doc.add_heading(f"{i}. {TEN_NHOM[nhom]}", level=1)
        if nhom == "V":
            bang(["Mã", "Đầu vào", "Kỳ vọng", "A", "B", "E(A)", "E(B)"],
                 [[a["ma"], a["dau_vao"], a["ky_vong"], dat_chu(a["dat"]), dat_chu(b and b["dat"]),
                   so(a["diem_E"]), so(b and b["diem_E"])] for a, b in cap_ca(bc, nhom)],
                 rong=[1.2, 3.8, 4.5, 2.2, 2.2, 1.1, 1.1], can_phai=(0, 3, 4, 5, 6))
        else:
            bang(["Mã", "Đầu vào / kịch bản", "Kỳ vọng", "A", "B"],
                 [[a["ma"], a["dau_vao"], a["ky_vong"], dat_chu(a["dat"]), dat_chu(b and b["dat"])]
                  for a, b in cap_ca(bc, nhom)], rong=[1.3, 5.5, 5.5, 2.2, 2.2], can_phai=(0, 3, 4))

    ls = bc["lich_su"]
    if ls["so_lan"] >= 2:
        doc.add_heading(f"6. Độ lặp lại qua {ls['so_lan']} lần chạy", level=1)
        bang(["Mốc", "Lặp", "F A/B", "V A/B", "S A/B", "E A/B", "T A/B (ms)"],
             [[x["moc"], x["lap"]] + [f"{so(x['A'][k], 3)}/{so(x['B'][k], 3)}"
                                      for k in ("F_%", "V_%", "S_%", "E_0_2", "T_ms")] for x in ls["cac_lan"]],
             can_phai=(1, 2, 3, 4, 5, 6))
        for pb in ("A", "B"):
            t = ls["thong_ke"][pb]
            doc.add_paragraph(f"Phiên bản {TEN_PB[pb]}: T trung bình {so(t['T_tb'], 3)} ms, độ lệch chuẩn "
                              f"{so(t['T_dlc'], 3)} ms; F, V, S, E "
                              + ("giống nhau ở mọi lần chạy." if t["FVSE_on_dinh"] else "thay đổi giữa các lần chạy."),
                              style="List Bullet")

    doc.add_heading("Phụ lục. Phản hồi thực tế của từng ca", level=1)
    bang(["PB", "Mã", "Kết quả", "T (ms)", "Phản hồi"],
         [[c["phien_ban"], c["ma"], dat_chu(c["dat"]), so(c["ms"], 3), c["phan_hoi"]] for c in bc["cac_ca"]],
         rong=[1, 1.2, 2.2, 1.8, 9.8], can_phai=(0, 1, 2, 3))
    p = doc.add_paragraph()
    r = p.add_run("Ghi chú: T đo bằng client trong bộ nhớ; giá trị tuyệt đối phụ thuộc máy, chỉ so sánh A với B "
                  f"trong cùng một lần chạy. Báo cáo xuất tự động lúc {bc['thoi_diem_xuat']}.")
    r.italic = True; r.font.size = Pt(11)
    doc.save(duong_dan)
    return True


# ============================================================ JSON & JUnit
def xuat_json(bc, duong_dan):
    with open(duong_dan, "w", encoding="utf-8") as f:
        json.dump(bc, f, ensure_ascii=False, indent=2)


def xuat_junit(bc, duong_dan, tinh_A):
    """Mỗi phiên bản là một testsuite; thêm suite 'Muc_tieu_Bang_10'.
    Ca KHÔNG ĐẠT của A (đối chứng) mặc định ghi là <skipped> để CI chỉ đỏ khi B hỏng; --junit-tinh-A để tính là failure."""
    goc = ET.Element("testsuites", name="Kiem thu MCP server – KHKT")
    tong = {"tests": 0, "failures": 0, "skipped": 0, "time": 0.0}

    def them_suite(ten, ca_list):
        s = ET.SubElement(goc, "testsuite", name=ten, timestamp=datetime.now().isoformat(timespec="seconds"))
        pr = ET.SubElement(s, "properties")
        for k, v in (("python", bc["moi_truong"]["python"]), ("mcp", bc["moi_truong"]["mcp"]),
                     ("lap", bc["moi_truong"]["lap"]), ("moc", bc["moc"])):
            ET.SubElement(pr, "property", name=k, value=str(v))
        dem = {"tests": 0, "failures": 0, "skipped": 0, "time": 0.0}
        for classname, name, t_giay, dat, loai_kd, thong_diep, chi_tiet in ca_list:
            tc = ET.SubElement(s, "testcase", classname=classname, name=name, time=f"{t_giay:.6f}")
            dem["tests"] += 1; dem["time"] += t_giay
            if not dat:
                if loai_kd == "skipped":
                    ET.SubElement(tc, "skipped", message=thong_diep); dem["skipped"] += 1
                else:
                    ET.SubElement(tc, "failure", message=thong_diep, type="KhongDat").text = chi_tiet
                    dem["failures"] += 1
        for k in ("tests", "failures", "skipped"):
            s.set(k, str(dem[k])); tong[k] += dem[k]
        s.set("errors", "0"); s.set("time", f"{dem['time']:.6f}"); tong["time"] += dem["time"]

    for pb in ("A", "B"):
        ds = []
        for c in [x for x in bc["cac_ca"] if x["phien_ban"] == pb]:
            loai = "skipped" if (pb == "A" and not tinh_A) else "failure"
            tm = ("Phiên bản đối chứng A không đạt (dự kiến theo thiết kế)" if loai == "skipped"
                  else f"Không đạt – kỳ vọng: {c['ky_vong']}")
            ds.append((f"{pb}.{c['nhom']}", f"{c['ma']} – {c['dau_vao']}", (c["ms"] or 0) / 1000, c["dat"], loai, tm,
                       f"Kỳ vọng: {c['ky_vong']}\nPhản hồi: {c['phan_hoi']}"))
        them_suite(f"Phien_ban_{'A_doi_chung' if pb == 'A' else 'B_cai_tien'}", ds)
    them_suite("Muc_tieu_Bang_10", [
        ("B.MucTieu", f"{m['chi_so']} – {m['muc_tieu']}", 0.0, m["dat"], "failure",
         f"B = {so(m['B'], 3)} chưa đạt mục tiêu {m['muc_tieu']}", f"A = {m['A']}, B = {m['B']}")
        for m in bc["muc_tieu"] if m["dat"] is not None])
    for k in ("tests", "failures", "skipped"):
        goc.set(k, str(tong[k]))
    goc.set("errors", "0"); goc.set("time", f"{tong['time']:.6f}")
    ET.indent(goc) if hasattr(ET, "indent") else None
    ET.ElementTree(goc).write(duong_dan, encoding="utf-8", xml_declaration=True)


# ============================================================ chương trình chính
def chay_kiem_thu(lap):
    lenh = [sys.executable, os.path.join(GOC, "run_tests.py"), "--lap", str(lap)]
    print(f">> Chạy 24 ca kiểm thử (lặp {lap} lần/ca)…")
    t0 = time.time()
    r = subprocess.run(lenh, cwd=GOC, env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    if r.returncode != 0:
        print(f"!! run_tests.py thất bại (mã {r.returncode}). Kiểm tra đã cài: pip install -r requirements.txt")
        sys.exit(2)
    print(f">> Xong trong {time.time() - t0:.1f} giây.\n")


def mo_tep(p):
    try:
        if sys.platform == "darwin":
            subprocess.run(["open", p], check=False)
        elif os.name == "nt":
            os.startfile(p)  # noqa
        else:
            subprocess.run(["xdg-open", p], check=False, stderr=subprocess.DEVNULL)
    except Exception as ex:  # không mở được cũng không sao
        print(f"  (không tự mở được: {ex})")


def main():
    ap = argparse.ArgumentParser(description="Xuất báo cáo kiểm thử MCP server ra HTML, DOCX, Markdown, JSON, JUnit XML.")
    ap.add_argument("--chay", action="store_true", help="chạy run_tests.py trước khi xuất báo cáo")
    ap.add_argument("--lap", type=int, default=20, help="số lần lặp mỗi ca khi --chay (mặc định 20)")
    ap.add_argument("--moc", help="mốc lần chạy cần xuất (mặc định: mới nhất)")
    ap.add_argument("--dinh-dang", default=",".join(DINH_DANG),
                    help=f"các định dạng, cách nhau dấu phẩy (mặc định: {','.join(DINH_DANG)})")
    ap.add_argument("--thu-muc-ra", help="thư mục ghi báo cáo (mặc định: ket_qua/bao_cao_<mốc>)")
    ap.add_argument("--nguong-T", type=float, default=1.5,
                    help="T của B được coi là 'không tăng đáng kể' nếu ≤ ngưỡng × T của A (mặc định 1,5)")
    ap.add_argument("--junit-tinh-A", action="store_true", help="tính ca không đạt của A là failure trong junit.xml")
    ap.add_argument("--mo", action="store_true", help="mở báo cáo HTML sau khi xuất")
    ap.add_argument("--liet-ke", action="store_true", help="liệt kê các lần chạy trong ket_qua/ rồi thoát")
    a = ap.parse_args()

    if a.liet_ke:
        ds = cac_moc()
        print("Chưa có lần chạy nào." if not ds else "\n".join(f"  {m}   ({moc_sang_ngay(m)})" for m in ds))
        return 0

    chon = [d.strip().lower() for d in a.dinh_dang.split(",") if d.strip()]
    sai = [d for d in chon if d not in DINH_DANG]
    if sai:
        print(f"!! Định dạng không hỗ trợ: {', '.join(sai)}. Hợp lệ: {', '.join(DINH_DANG)}")
        return 2

    if a.chay:
        chay_kiem_thu(a.lap)
    ds = cac_moc()
    if not ds:
        print("!! Chưa có kết quả trong ket_qua/. Chạy:  python xuat_bao_cao.py --chay")
        return 2
    moc = a.moc or ds[-1]
    if moc not in ds:
        print(f"!! Không có lần chạy {moc}. Dùng --liet-ke để xem các mốc.")
        return 2

    try:
        bc = dung_bao_cao(moc, a.nguong_T)
    except (FileNotFoundError, KeyError, ValueError) as ex:
        print(f"!! Dữ liệu lần chạy {moc} không đầy đủ: {ex}")
        return 2

    ra = a.thu_muc_ra or os.path.join(KET_QUA, f"bao_cao_{moc}")
    os.makedirs(ra, exist_ok=True)
    print(f">> Xuất báo cáo lần chạy {moc_sang_ngay(moc)} vào {os.path.relpath(ra, GOC)}/")
    tep = {"html": "bao_cao.html", "docx": "bao_cao.docx", "md": "bao_cao.md", "json": "ket_qua.json", "junit": "junit.xml"}
    for d in chon:
        p = os.path.join(ra, tep[d])
        ok = {"html": lambda: xuat_html(bc, p), "docx": lambda: xuat_docx(bc, p), "md": lambda: xuat_md(bc, p),
              "json": lambda: xuat_json(bc, p), "junit": lambda: xuat_junit(bc, p, a.junit_tinh_A)}[d]()
        if ok is not False:
            print(f"   ✓ {tep[d]}")

    print()
    for c in bc["ket_luan"]:
        print(" •", c)
    print("\nKẾT QUẢ:", "B ĐẠT mọi mục tiêu Bảng 10." if bc["B_dat_moi_muc_tieu"] else "B CHƯA ĐẠT đủ mục tiêu Bảng 10.")
    if a.mo and "html" in chon:
        mo_tep(os.path.join(ra, "bao_cao.html"))
    return 0 if bc["B_dat_moi_muc_tieu"] else 1


if __name__ == "__main__":
    sys.exit(main())
