"""Chạy kiểm thử và TỰ ĐỘNG GHI NHẬT KÝ NGHIÊN CỨU theo mẫu Phụ lục 2 (Hướng dẫn sổ nhật ký nghiên cứu).

Mỗi lần chạy = MỘT trang nhật ký mới, được NỐI THÊM vào cuối sổ (không sửa, không xoá trang cũ):
  nhat_ky_nghien_cuu/So_nhat_ky_nghien_cuu.docx   – bản in (bìa + mục lục + các trang nhật ký)
  nhat_ky_nghien_cuu/So_nhat_ky_nghien_cuu.md     – bản văn bản thuần, cùng nội dung
  nhat_ky_nghien_cuu/trang/<số trang>_<mốc>.json   – toàn bộ số liệu của trang (để đối chiếu)
  nhat_ky_nghien_cuu/trang_thai.json              – số trang + chuỗi mã băm SHA-256 nối các trang
                                                    (trang sau chứa mã băm trang trước -> phát hiện được
                                                    nếu có trang bị sửa/xoá, tương đương quy định "không xé trang")

Trình tự mỗi lần chạy:
  1. Sinh bộ dữ liệu mô phỏng nếu chưa có (sinh_du_lieu_kiem_thu.py, seed cố định).
  2. Chạy bộ 24 ca chuẩn (run_tests.py) --lan lần  -> "Lần thử 1, 2, 3".
  3. Chạy bộ ca mở rộng (chay_kiem_thu_mo_rong.py) --lan lần.
  4. Ghi trang nhật ký: NGÀY / TRANG, GIAI ĐOẠN, THỜI GIAN; 1. Mục tiêu; 2. Dụng cụ & vật liệu;
     3. Tiến trình & hiện tượng (kèm sơ đồ); 4. Kết quả & số liệu thô (lần 1, 2, 3, trung bình);
     5. Rút kinh nghiệm & lỗi sai; 6. Kế hoạch tiếp theo; chữ ký học sinh / giáo viên.

LƯU Ý QUAN TRỌNG: Phụ lục 2 yêu cầu sổ nhật ký chính VIẾT TAY BẰNG BÚT MỰC. Trang in tự động này là
PHỤ LỤC SỐ LIỆU THÔ, dán/kẹp vào sổ viết tay; học sinh vẫn tự ghi nhận xét, rút kinh nghiệm và ký tên.

Chạy:   python ghi_nhat_ky.py
        python ghi_nhat_ky.py --lan 3 --lap 20 --giai-doan "Thực nghiệm lần 2" --ghi-chu "Chạy trên máy phòng Tin"
        python ghi_nhat_ky.py --kiem-tra        (chỉ kiểm tra chuỗi mã băm, không chạy)
Cần: mcp==2.2.0, python-docx (pip install -r requirements.txt -r requirements-bao-cao.txt)
"""

import argparse
import asyncio
import csv
import glob
import hashlib
import json
import os
import platform
import statistics
import subprocess
import sys
import time
from datetime import datetime
from importlib.metadata import version

# Sổ ghi theo giờ Việt Nam (UTC+7) kể cả khi máy/VM đặt múi giờ UTC hoặc thiếu dữ liệu múi giờ;
# tiến trình con (run_tests.py) thừa hưởng biến TZ này.
if hasattr(time, "tzset") and time.localtime().tm_gmtoff == 0:
    os.environ["TZ"] = "ICT-7"  # chuỗi POSIX = UTC+7, không cần cơ sở dữ liệu múi giờ
    time.tzset()

GOC = os.path.dirname(os.path.abspath(__file__))
THU_MUC_NK = os.path.join(GOC, "nhat_ky_nghien_cuu")
TEP_DOCX = os.path.join(THU_MUC_NK, "So_nhat_ky_nghien_cuu.docx")
TEP_MD = os.path.join(THU_MUC_NK, "So_nhat_ky_nghien_cuu.md")
TEP_TT = os.path.join(THU_MUC_NK, "trang_thai.json")
DE_TAI = ("Thiết kế và thử nghiệm chương trình giúp học sinh phổ thông kết nối trí tuệ nhân tạo "
          "với công cụ bên ngoài thông qua giao thức MCP")
LINH_VUC = "Phần mềm hệ thống (chuyên sâu: An ninh máy tính; Cơ sở dữ liệu)"
# Bản công khai: thông tin học sinh (họ tên, ngày sinh) được để trống để bảo vệ quyền riêng tư.
# Điền thông tin thật vào bản chạy trên máy của nhóm trước khi in sổ nhật ký.
HOC_SINH = [("<Học sinh 1>", "nhóm trưởng"),
            ("<Học sinh 2>", "")]
GIAO_VIEN_HD = "Đặng Việt Chung"
TRUONG = "Trường THPT Phước Bửu"
CHI_SO = [("F_%", "F – chức năng (%)"), ("V_%", "V – từ chối đầu vào bất thường (%)"),
          ("S_%", "S – an toàn (%)"), ("E_0_2", "E – chất lượng thông báo lỗi (0–2)"),
          ("T_ms", "T – thời gian phản hồi TB (ms)"), ("dat", "Số ca đạt / 24")]


# ------------------------------------------------------------------ tiện ích
def sha(du_lieu: bytes) -> str:
    return hashlib.sha256(du_lieu).hexdigest()


def sha_tep(p):
    return sha(open(p, "rb").read()) if os.path.exists(p) else None


def so(x, n=2):
    """Định dạng số kiểu Việt Nam (dấu phẩy thập phân)."""
    if isinstance(x, (int, float)):
        return f"{x:.{n}f}".rstrip("0").rstrip(".").replace(".", ",") if isinstance(x, float) else str(x)
    return str(x)


def doc_tt():
    if os.path.exists(TEP_TT):
        return json.load(open(TEP_TT, encoding="utf-8"))
    return {"so_trang": 0, "cac_trang": [], "sha_docx": None}


def kiem_tra_chuoi(tt):
    """Kiểm tra: mỗi trang còn nguyên (mã băm khớp) và nối đúng trang trước. Trả về danh sách vấn đề."""
    van_de, truoc = [], None
    for t in tt["cac_trang"]:
        p = os.path.join(THU_MUC_NK, t["tep"])
        if not os.path.exists(p):
            van_de.append(f"Trang {t['trang']}: thiếu tệp {t['tep']}")
        elif sha_tep(p) != t["sha256"]:
            van_de.append(f"Trang {t['trang']}: nội dung đã bị thay đổi sau khi ghi")
        if t["sha_trang_truoc"] != truoc:
            van_de.append(f"Trang {t['trang']}: không nối đúng trang trước")
        truoc = t["sha256"]
    if tt.get("sha_docx") and os.path.exists(TEP_DOCX) and sha_tep(TEP_DOCX) != tt["sha_docx"]:
        van_de.append("Tệp .docx đã được chỉnh sửa ngoài công cụ kể từ lần ghi trước (trang mới vẫn được nối thêm).")
    return van_de


# ------------------------------------------------------------------ chạy kiểm thử
def chay_24_ca(lan, lap):
    kq = []
    for i in range(1, lan + 1):
        truoc = set(glob.glob(os.path.join(GOC, "ket_qua", "tong_hop_*.json")))
        bd = datetime.now()
        r = subprocess.run([sys.executable, "run_tests.py", "--lap", str(lap)], cwd=GOC,
                           capture_output=True, text=True, encoding="utf-8")
        moi = sorted(set(glob.glob(os.path.join(GOC, "ket_qua", "tong_hop_*.json"))) - truoc)
        if r.returncode != 0 or not moi:
            raise SystemExit(f"run_tests.py lỗi ở lần thử {i}:\n{r.stdout[-2000:]}\n{r.stderr[-2000:]}")
        th = json.load(open(moi[-1], encoding="utf-8"))
        moc = th["thoi_diem"]
        ca_truot = {"A": [], "B": []}
        with open(os.path.join(GOC, "ket_qua", f"so_lieu_tho_{moc}.csv"), encoding="utf-8-sig") as f:
            for d in csv.DictReader(f):
                if d["dat"] not in ("True", "1"):
                    ca_truot[d["phien_ban"]].append(d["ma"])
        kq.append({"lan": i, "moc": moc, "bat_dau": bd.strftime("%H:%M:%S"),
                   "ket_thuc": datetime.now().strftime("%H:%M:%S"), "chi_so": th["chi_so"],
                   "ca_khong_dat": ca_truot,
                   "tep": [f"ket_qua/so_lieu_tho_{moc}.csv", f"ket_qua/tong_hop_{moc}.json",
                           f"ket_qua/bang_ket_qua_{moc}.md"]})
        time.sleep(1.1)  # bảo đảm mốc thời gian (đến giây) của các lần thử khác nhau
    return kq


def chay_mo_rong(lan):
    import chay_kiem_thu_mo_rong as m  # nhập muộn để --kiem-tra không cần thư viện mcp
    return asyncio.run(m.chay(lan))


def cac_lan_chay_trong_ngay(ngay_yyyymmdd, moc_da_dung):
    """Liệt kê mọi lần chạy bộ 24 ca trong ngày (kể cả lần do học sinh tự chạy), trừ các mốc của trang này."""
    ds = []
    for f in sorted(glob.glob(os.path.join(GOC, "ket_qua", f"tong_hop_{ngay_yyyymmdd}_*.json"))):
        th = json.load(open(f, encoding="utf-8"))
        if th["thoi_diem"] in moc_da_dung:
            continue
        c = th["chi_so"]
        gio = th["thoi_diem"][9:11] + ":" + th["thoi_diem"][11:13]
        ds.append(f"{gio} – lặp {th.get('lap')} lần/ca, Python {th.get('python')}: "
                  f"A {c['A']['dat']} (V {c['A']['V_%']}%, S {c['A']['S_%']}%, T {c['A']['T_ms']} ms); "
                  f"B {c['B']['dat']} (E {c['B']['E_0_2']}, T {c['B']['T_ms']} ms)")
    return ds


def trung_binh(ds):
    so_ = [x for x in ds if isinstance(x, (int, float))]
    return round(statistics.mean(so_), 3) if so_ else None


def hien_tuong(dong_mr, bo_ca):
    """Rút ra các hiện tượng đáng chú ý từ số liệu thô lần thử 1 của bộ mở rộng (tự động, khách quan)."""
    tra = {c["ma"]: c for c in bo_ca["ca"]}
    d1 = [x for x in dong_mr if x["lan"] == 1]
    out = []
    A_nhan_sai = [x for x in d1 if x["ver"] == "A" and x["ky_vong"] == "tu_choi" and not x["bi_tu_choi"]
                  and x["nhom"].startswith("MV")]
    if A_nhan_sai:
        vd = "; ".join(f"{tra[x['ma']]['mo_ta']} -> “{x['phan_hoi_rut_gon'][:50] or '(rỗng/None)'}”"
                       for x in A_nhan_sai[:4])
        out.append(f"A CHẤP NHẬN {len(A_nhan_sai)} đầu vào sai mà không báo lỗi. Ví dụ: {vd}.")
    ro = [x for x in d1 if x["ver"] == "A" and x["ro_ri"]]
    if ro:
        out.append(f"A LÀM LỘ mã kiểm tra ở {len(ro)} ca truy cập tệp: "
                   + ", ".join(f"{x['ma']} ({tra[x['ma']]['mo_ta'].replace('Truy cập tệp: ', '')})" for x in ro) + ".")
    e0 = [x for x in d1 if x["ver"] == "A" and x["E"] == 0 and x["bi_tu_choi"]]
    if e0:
        out.append(f"A từ chối {len(e0)} ca nhưng thông báo lỗi chung chung (E = 0), ví dụ: "
                   f"“{e0[0]['phan_hoi_rut_gon'][:60]}”.")
    b_loi = [x for x in d1 if x["ver"] == "B" and x["bi_tu_choi"] and x["E"] is not None]
    if b_loi:
        vd = next((x for x in b_loi if x["E"] == 2), b_loi[0])
        out.append(f"B từ chối {len(b_loi)} ca đầu vào sai; ví dụ thông báo ({vd['ma']}): "
                   f"“{vd['phan_hoi_rut_gon'][:110]}”.")
    b_truot = [x for x in d1 if x["ver"] == "B" and not x["dat"]]
    out.append("B không có ca nào trượt ở lần thử 1." if not b_truot else
               "B TRƯỢT các ca: " + ", ".join(f"{x['ma']} ({x['ghi_chu']})" for x in b_truot) + ".")
    tiem = [x for x in d1 if x["nhom"] == "MS_TIEM"]
    for x in tiem:
        out.append(f"Ca chèn lệnh qua ghi chú thứ Năm – {x['ver']}: {x['ghi_chu']}.")
    return out


# ------------------------------------------------------------------ nội dung trang
def lap_trang(a, tt, bat_dau, kq24, th_mr, dong_mr, bo_ca, van_de):
    ket_thuc = datetime.now()
    trang = tt["so_trang"] + 1
    tt_sinh = json.load(open(os.path.join(GOC, "du_lieu", "kiem_thu", "thong_tin_sinh.json"), encoding="utf-8"))
    tep_ma = ["server_a.py", "server_b.py", "run_tests.py", "sinh_du_lieu_kiem_thu.py", "chay_kiem_thu_mo_rong.py",
              "du_lieu/thoi_khoa_bieu.json", "du_lieu/lich_thi.json", "du_lieu/noi_quy.txt",
              "du_lieu/kiem_thu/bo_ca_mo_rong.json", "du_lieu/kiem_thu/diem_lop_mo_phong.csv"]
    dung_cu = [
        f"Máy tính: {platform.system()} {platform.release()} ({platform.machine()}); Python {platform.python_version()}; "
        f"thư viện mcp {version('mcp')} (MCP Python SDK); client kết nối trong bộ nhớ, không mở cổng mạng.",
        "Phần mềm thử nghiệm: server_a.py (phiên bản A – đối chứng), server_b.py (phiên bản B – cải tiến).",
        "Công cụ đo: run_tests.py (24 ca chuẩn F1–F10, V1–V8, S1–S6; lặp "
        f"{a.lap} lần/ca để đo T), chay_kiem_thu_mo_rong.py (bộ ca mở rộng).",
        f"Vật liệu: dữ liệu giả lập lớp 11A (thời khoá biểu, lịch thi, nội quy) + bộ mô phỏng {tt_sinh['so_ca']} ca "
        f"sinh từ hạt giống {tt_sinh['seed']} ({tt_sinh['so_hoc_sinh_gia_lap']} học sinh giả lập HS001…, "
        "không có dữ liệu thật); 2 tệp mồi vô hại chứa mã [MA_KIEM_TRA_S1], [MA_KIEM_TRA_S2].",
    ]
    ma_bam = [(t, (sha_tep(os.path.join(GOC, t)) or "không có")[:16]) for t in tep_ma]

    tien_trinh = [f"{bat_dau:%H:%M:%S} – Bắt đầu; kiểm tra dữ liệu mô phỏng (seed {tt_sinh['seed']}, "
                  f"{tt_sinh['so_ca']} ca, sinh lúc {tt_sinh['thoi_diem_sinh']})."]
    for k in kq24:
        tien_trinh.append(f"{k['bat_dau']}–{k['ket_thuc']} – Lần thử {k['lan']}: chạy 24 ca chuẩn trên A và B "
                          f"(mốc {k['moc']}).")
    tien_trinh.append(f"{th_mr['moc'][9:11]}:{th_mr['moc'][11:13]}:{th_mr['moc'][13:15]} – Chạy bộ ca mở rộng "
                      f"{th_mr['so_lan']} lần × {th_mr['so_ca']} ca × 2 phiên bản = {th_mr['so_lan'] * th_mr['so_ca'] * 2} lượt gọi.")
    tien_trinh.append(f"{ket_thuc:%H:%M:%S} – Tổng hợp số liệu, ghi trang nhật ký.")

    # bảng 24 ca: chỉ số × (A lần 1..n, TB, B lần 1..n, TB)
    bang24 = []
    for khoa, ten in CHI_SO:
        hang = [ten]
        for ver in ("A", "B"):
            gt = [k["chi_so"][ver][khoa] for k in kq24]
            hang += [so(x) for x in gt] + [so(trung_binh(gt)) if khoa != "dat" else "–"]
        bang24.append(hang)
    # bảng mở rộng: nhóm × (A lần.., B lần..) tỉ lệ đạt
    nhom_ds = sorted(th_mr["gop"]["A"]["theo_nhom"])
    bang_mr = []
    for n in nhom_ds:
        hang = [f"{n} ({th_mr['gop']['A']['theo_nhom'][n]['so_ca'] // th_mr['so_lan']} ca)"]
        for ver in ("A", "B"):
            gt = [l[ver]["theo_nhom"][n]["ti_le_%"] for l in th_mr["theo_lan"]]
            hang += [so(x, 1) for x in gt] + [so(trung_binh(gt), 1)]
        bang_mr.append(hang)
    for nhan, khoa, nd in (("Tỉ lệ đạt chung (%)", "ti_le_dat_%", 1), ("E trung bình (0–2)", "E_tb_0_2", 2),
                           ("Số lần lộ mã kiểm tra", "so_lan_ro_ri", 0), ("T trung bình (ms)", "T_tb_ms", 3)):
        hang = [nhan]
        for ver in ("A", "B"):
            gt = [l[ver][khoa] for l in th_mr["theo_lan"]]
            hang += [so(x, nd) for x in gt] + [so(trung_binh(gt), nd)]
        bang_mr.append(hang)

    # rút kinh nghiệm tự động
    rkn = []
    for ver in ("A", "B"):
        T = [k["chi_so"][ver]["T_ms"] for k in kq24]
        if len(T) > 1 and statistics.mean(T) > 0:
            cv = 100 * statistics.pstdev(T) / statistics.mean(T)
            rkn.append(f"T của {ver} giữa các lần thử dao động {so(cv, 1)}% (hệ số biến thiên)"
                       + (" – lớn, nên đóng bớt chương trình khác và chạy lại." if cv > 20 else " – chấp nhận được."))
        dat = {k["chi_so"][ver]["dat"] for k in kq24}
        if len(dat) > 1:
            rkn.append(f"Số ca đạt của {ver} KHÁC NHAU giữa các lần thử ({', '.join(sorted(dat))}) – cần tìm nguyên nhân.")
    ca_B = sorted({m for k in kq24 for m in k["ca_khong_dat"]["B"]})
    rkn.append("B đạt đủ 24/24 ca ở mọi lần thử." if not ca_B else f"B trượt ca: {', '.join(ca_B)} – cần sửa.")
    ca_A = sorted({m for k in kq24 for m in k["ca_khong_dat"]["A"]})
    rkn.append(f"A (đối chứng, cố ý viết thiếu nguyên tắc) trượt ca: {', '.join(ca_A) or 'không có'} – đúng như thiết kế thí nghiệm.")
    rkn += [f"Cảnh báo tính toàn vẹn sổ: {v}" for v in van_de]
    rkn.append("Số liệu thô được giữ nguyên trong các tệp CSV/JSON đã liệt kê, kể cả các ca lỗi.")

    ke_hoach = []
    if ca_B:
        ke_hoach.append("Sửa server_b.py cho các ca B còn trượt, chạy lại toàn bộ và ghi trang mới.")
    ke_hoach += ["Đối chiếu Bảng 7–10 trong báo cáo với trung bình 3 lần thử ở trang này.",
                 "Thử lại với hạt giống khác (python sinh_du_lieu_kiem_thu.py --seed <số>) để kiểm tra độ ổn định.",
                 "Kiểm tra thủ công bằng MCP Inspector một vài ca tiêu biểu và ghi tay nhận xét."]

    return {
        "trang": trang, "ngay": ket_thuc.strftime("%d/%m/%Y"), "giai_doan": a.giai_doan,
        "tu": bat_dau.strftime("%H:%M"), "den": ket_thuc.strftime("%H:%M"),
        "muc_tieu": a.muc_tieu or [
            "Chạy lại bộ 24 ca kiểm thử chuẩn trên hai phiên bản A và B, lặp "
            f"{len(kq24)} lần thử để kiểm tra độ lặp lại của F, V, S, E, T.",
            f"Chạy bộ dữ liệu mô phỏng mở rộng {th_mr['so_ca']} ca để kiểm tra kết luận trên nhiều đầu vào hơn.",
            "Ghi nhận toàn bộ số liệu thô, kể cả ca lỗi."],
        "dung_cu": dung_cu, "ma_bam": ma_bam, "tien_trinh": tien_trinh,
        "hien_tuong": hien_tuong(dong_mr, bo_ca),
        "so_lan": len(kq24), "bang24": bang24, "bang_mr": bang_mr,
        "tep_so_lieu": [t for k in kq24 for t in k["tep"]] + list(th_mr["tep"].values()),
        "lan_chay_khac": cac_lan_chay_trong_ngay(ket_thuc.strftime("%Y%m%d"), {k["moc"] for k in kq24}),
        "hoc_sinh": [f"{ten} ({gc})" for ten, gc in HOC_SINH], "giao_vien_hd": GIAO_VIEN_HD,
        "rut_kinh_nghiem": rkn, "ke_hoach": ke_hoach, "ghi_chu_hs": a.ghi_chu,
        "sha_trang_truoc": tt["cac_trang"][-1]["sha256"] if tt["cac_trang"] else None,
        "kq24": kq24, "mo_rong": {k: th_mr[k] for k in ("moc", "seed", "so_ca", "so_lan", "gop")},
    }


# ------------------------------------------------------------------ xuất Markdown
def md_trang(p):
    L = [f"\n---\n\n## NGÀY: {p['ngay']}  |  TRANG: {p['trang']}",
         f"**GIAI ĐOẠN:** {p['giai_doan']}  |  **THỜI GIAN:** từ {p['tu']} đến {p['den']}",
         f"**Người thực hiện:** {'; '.join(p.get('hoc_sinh', []))}  |  **GVHD:** {p.get('giao_vien_hd','')}", "",
         "### 1. MỤC TIÊU"] + [f"- {x}" for x in p["muc_tieu"]] + ["", "### 2. DỤNG CỤ & VẬT LIỆU"]
    L += [f"- {x}" for x in p["dung_cu"]]
    L += ["- Mã băm SHA-256 (16 ký tự đầu) của mã nguồn và dữ liệu dùng trong lần chạy:"]
    L += [f"  - `{t}`: `{h}`" for t, h in p["ma_bam"]]
    L += ["", "### 3. TIẾN TRÌNH & HIỆN TƯỢNG",
          "Sơ đồ: `Bộ ca (seed) → Client MCP (trong bộ nhớ) → Server A | Server B → Chấm tự động → CSV/JSON → Sổ nhật ký`",
          "", "**Tiến trình:**"] + [f"- {x}" for x in p["tien_trinh"]]
    L += ["", "**Hiện tượng quan sát được (tự động trích từ số liệu thô, lần thử 1):**"] + [f"- {x}" for x in p["hien_tuong"]]
    n = p["so_lan"]
    tieu_de = ["Chỉ số"] + [f"A lần {i}" for i in range(1, n + 1)] + ["A TB"] + [f"B lần {i}" for i in range(1, n + 1)] + ["B TB"]
    L += ["", "### 4. KẾT QUẢ & SỐ LIỆU THÔ", "", "**a) Bộ 24 ca chuẩn**", "",
          "| " + " | ".join(tieu_de) + " |", "|" + "---|" * len(tieu_de)]
    L += ["| " + " | ".join(h) + " |" for h in p["bang24"]]
    tieu_de[0] = "Nhóm ca mở rộng (tỉ lệ đạt %)"
    L += ["", "**b) Bộ ca mô phỏng mở rộng**", "", "| " + " | ".join(tieu_de) + " |", "|" + "---|" * len(tieu_de)]
    L += ["| " + " | ".join(h) + " |" for h in p["bang_mr"]]
    if p.get("lan_chay_khac"):
        L += ["", "**c) Các lần chạy khác trong ngày (ghi nhận nguyên trạng):**"] + [f"- {x}" for x in p["lan_chay_khac"]]
    L += ["", "**Tệp số liệu thô (giữ nguyên, kể cả ca lỗi):**"] + [f"- `{t}`" for t in p["tep_so_lieu"]]
    L += ["", "### 5. RÚT KINH NGHIỆM & LỖI SAI"] + [f"- {x}" for x in p["rut_kinh_nghiem"]]
    if p["ghi_chu_hs"]:
        L += [f"- Ghi chú của học sinh: {p['ghi_chu_hs']}"]
    L += ["- Học sinh tự ghi thêm: ....................................................................",
          "", "### 6. KẾ HOẠCH TIẾP THEO"] + [f"- {x}" for x in p["ke_hoach"]]
    L += ["- Học sinh tự ghi thêm: ....................................................................", "",
          f"| Chữ ký học sinh ({' / '.join(t for t, _ in HOC_SINH)}) | Chữ ký giáo viên hướng dẫn ({GIAO_VIEN_HD}) |",
         "|---|---|", "| &nbsp;<br><br> | &nbsp;<br><br> |", "",
          f"<sub>Trang in tự động – phụ lục số liệu thô kèm sổ nhật ký viết tay. Mã băm trang trước: "
          f"{(p['sha_trang_truoc'] or 'không có (trang đầu)')[:16]}</sub>"]
    return "\n".join(L) + "\n"


# ------------------------------------------------------------------ xuất DOCX
def _docx():
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    return Document, WD_ALIGN_PARAGRAPH, WD_BREAK, qn, Cm, Pt


def _font(doc, qn, Pt):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def tao_so_moi(ngay):
    Document, CENTER, WD_BREAK, qn, Cm, Pt = _docx()
    CENTER = CENTER.CENTER
    doc = Document()
    _font(doc, qn, Pt)
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
    for t, c, b in (("SỞ GIÁO DỤC VÀ ĐÀO TẠO", 13, False), ("TRƯỜNG THPT ………………………………", 13, True),
                    ("", 13, False), ("", 13, False), ("SỔ NHẬT KÝ NGHIÊN CỨU", 26, True),
                    ("PHẦN IN TỰ ĐỘNG – PHỤ LỤC SỐ LIỆU THÔ KIỂM THỬ", 14, True), ("", 13, False)):
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run(t); r.bold = b; r.font.size = Pt(c)
    p = doc.add_paragraph(); p.alignment = CENTER
    r = p.add_run(f"Tên dự án: {DE_TAI}"); r.italic = True
    dong_bia = ["", f"Lĩnh vực dự thi: {LINH_VUC}"]
    dong_bia += [f"Học sinh thực hiện: {ten} ({gc}) – Lớp: ……" for ten, gc in HOC_SINH]
    dong_bia += [f"Giáo viên hướng dẫn: {GIAO_VIEN_HD}", f"Đơn vị: {TRUONG}",
                 f"Ngày bắt đầu: {ngay}", "Ngày kết thúc: ……………"]
    for t in dong_bia:
        doc.add_paragraph(t)
    p = doc.add_paragraph(); p.alignment = CENTER
    r = p.add_run("\nSổ nhật ký chính được VIẾT TAY bằng bút mực theo Phụ lục 2. Các trang in tự động dưới đây "
                  "là phụ lục số liệu thô, được nối thêm theo thứ tự thời gian, không sửa, không xoá trang cũ.")
    r.italic = True; r.font.size = Pt(11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h = doc.add_paragraph(); h.alignment = CENTER
    r = h.add_run("MỤC LỤC"); r.bold = True; r.font.size = Pt(16)
    doc.add_paragraph("(Tự động điền khi thêm trang; học sinh có thể ghi tay bổ sung.)").runs[0].italic = True
    tb = doc.add_table(rows=1, cols=4); tb.style = "Table Grid"
    for i, t in enumerate(("Trang", "Ngày", "Giai đoạn", "Nội dung chính")):
        tb.rows[0].cells[i].text = t
        tb.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for _ in range(30):
        tb.add_row()
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc


def them_trang_docx(p):
    Document, ALIGN, WD_BREAK, qn, Cm, Pt = _docx()
    doc = Document(TEP_DOCX) if os.path.exists(TEP_DOCX) else tao_so_moi(p["ngay"])
    # --- mục lục: điền dòng trống đầu tiên (thêm dòng nếu đã đầy)
    ml = next((t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip() == "Trang"), None)
    if ml is not None:
        hang = next((r for r in ml.rows[1:] if not r.cells[0].text.strip()), None) or ml.add_row()
        for i, t in enumerate((str(p["trang"]), p["ngay"], p["giai_doan"],
                               f"Kiểm thử A/B: 24 ca × {p['so_lan']} lần + {p['mo_rong']['so_ca']} ca mô phỏng")):
            hang.cells[i].text = t
    if p["trang"] > 1:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def dong(nhan, gt=""):
        q = doc.add_paragraph(); r = q.add_run(nhan); r.bold = True; q.add_run(gt)
        q.paragraph_format.space_after = Pt(2)
        return q

    def muc(t):
        q = doc.add_paragraph(); r = q.add_run(t); r.bold = True; r.font.size = Pt(13.5)
        q.paragraph_format.space_before = Pt(8); q.paragraph_format.space_after = Pt(2)

    def gach(ds, co=12):
        for x in ds:
            q = doc.add_paragraph(style="List Bullet"); r = q.add_run(x); r.font.size = Pt(co)
            q.paragraph_format.space_after = Pt(0)

    def bang(tieu_de, hang, co=10):
        t = doc.add_table(rows=1, cols=len(tieu_de)); t.style = "Table Grid"
        for i, x in enumerate(tieu_de):
            c = t.rows[0].cells[i]; c.text = x; c.paragraphs[0].runs[0].bold = True
            c.paragraphs[0].runs[0].font.size = Pt(co)
        for h in hang:
            cells = t.add_row().cells
            for i, x in enumerate(h):
                cells[i].text = str(x); cells[i].paragraphs[0].runs[0].font.size = Pt(co)
        if len(tieu_de) > 2:  # cột đầu rộng hơn cho dễ đọc
            t.autofit = False
            for r in t.rows:
                r.cells[0].width = Cm(4.2)
                for c in r.cells[1:]:
                    c.width = Cm(11.8 / (len(tieu_de) - 1))
        return t

    tb = doc.add_table(rows=3, cols=2); tb.style = "Table Grid"
    tb.cell(2, 0).merge(tb.cell(2, 1))
    tb.cell(2, 0).text = ("Người thực hiện: " + "; ".join(p.get("hoc_sinh", []))
                          + "  |  GVHD: " + p.get("giao_vien_hd", ""))
    tb.cell(0, 0).text = f"NGÀY: {p['ngay']}"; tb.cell(0, 1).text = f"TRANG: {p['trang']}"
    tb.cell(1, 0).text = f"GIAI ĐOẠN: {p['giai_doan']}"; tb.cell(1, 1).text = f"THỜI GIAN: từ {p['tu']} đến {p['den']}"
    for c in (tb.cell(0, 0), tb.cell(0, 1), tb.cell(1, 0), tb.cell(1, 1)):
        c.paragraphs[0].runs[0].bold = True

    muc("1. MỤC TIÊU"); gach(p["muc_tieu"])
    muc("2. DỤNG CỤ & VẬT LIỆU"); gach(p["dung_cu"])
    dong("Mã băm SHA-256 (16 ký tự đầu) – để đối chiếu đúng phiên bản mã/dữ liệu:").runs[0].font.size = Pt(11)
    bang(["Tệp", "SHA-256"], p["ma_bam"], co=9)
    muc("3. TIẾN TRÌNH & HIỆN TƯỢNG")
    so_do = doc.add_table(rows=1, cols=9); so_do.style = "Table Grid"
    for i, x in enumerate(["Bộ ca\n(seed)", "→", "Client MCP\n(trong bộ nhớ)", "→", "Server A\nServer B", "→",
                           "Chấm tự động", "→", "CSV/JSON\n→ Sổ"]):
        c = so_do.rows[0].cells[i]; c.text = x
        c.paragraphs[0].alignment = ALIGN.CENTER; c.paragraphs[0].runs[0].font.size = Pt(9)
    dong("Tiến trình:"); gach(p["tien_trinh"], 11)
    dong("Hiện tượng quan sát được (tự động trích từ số liệu thô, lần thử 1):"); gach(p["hien_tuong"], 11)
    muc("4. KẾT QUẢ & SỐ LIỆU THÔ")
    n = p["so_lan"]
    td = ["Chỉ số"] + [f"A-L{i}" for i in range(1, n + 1)] + ["A TB"] + [f"B-L{i}" for i in range(1, n + 1)] + ["B TB"]
    dong("a) Bộ 24 ca chuẩn (L = lần thử, TB = trung bình):"); bang(td, p["bang24"], co=9)
    td[0] = "Nhóm ca mở rộng (tỉ lệ đạt %)"
    dong("b) Bộ ca mô phỏng mở rộng:"); bang(td, p["bang_mr"], co=9)
    if p.get("lan_chay_khac"):
        dong("c) Các lần chạy khác trong ngày (ghi nhận nguyên trạng):"); gach(p["lan_chay_khac"], 10)
    dong("Tệp số liệu thô (giữ nguyên, kể cả ca lỗi):"); gach(p["tep_so_lieu"], 10)
    muc("5. RÚT KINH NGHIỆM & LỖI SAI"); gach(p["rut_kinh_nghiem"], 11)
    if p["ghi_chu_hs"]:
        gach([f"Ghi chú của học sinh: {p['ghi_chu_hs']}"], 11)
    for _ in range(3):
        doc.add_paragraph("…" * 60)
    muc("6. KẾ HOẠCH TIẾP THEO"); gach(p["ke_hoach"], 11)
    for _ in range(2):
        doc.add_paragraph("…" * 60)
    ck = doc.add_table(rows=2, cols=2)
    ck.cell(0, 0).text = "Chữ ký học sinh\n" + " / ".join(t for t, _ in HOC_SINH)
    ck.cell(0, 1).text = "Chữ ký giáo viên hướng dẫn\n" + GIAO_VIEN_HD
    for c in (ck.cell(0, 0), ck.cell(0, 1)):
        c.paragraphs[0].alignment = ALIGN.CENTER; c.paragraphs[0].runs[0].bold = True
    ck.cell(1, 0).text = "\n\n\n"; ck.cell(1, 1).text = "\n\n\n"
    q = doc.add_paragraph()
    r = q.add_run(f"Trang in tự động – phụ lục số liệu thô kèm sổ nhật ký viết tay. Mã băm trang trước: "
                  f"{(p['sha_trang_truoc'] or 'không có (trang đầu)')[:16]}")
    r.italic = True; r.font.size = Pt(9)
    doc.save(TEP_DOCX)


def them_trang_md(p):
    moi = not os.path.exists(TEP_MD)
    with open(TEP_MD, "a", encoding="utf-8") as f:
        if moi:
            f.write(f"# SỔ NHẬT KÝ NGHIÊN CỨU – PHẦN IN TỰ ĐỘNG (PHỤ LỤC SỐ LIỆU THÔ)\n\n**Dự án:** {DE_TAI}\n\n"
                    "_Sổ chính viết tay bằng bút mực theo Phụ lục 2; các trang dưới đây chỉ được nối thêm, "
                    "không sửa, không xoá._\n")
        f.write(md_trang(p))


# ------------------------------------------------------------------ chính
def main():
    ap = argparse.ArgumentParser(description="Chạy kiểm thử và tự động ghi sổ nhật ký nghiên cứu (Phụ lục 2)")
    ap.add_argument("--lan", type=int, default=3, help="số lần thử (mặc định 3: Lần thử 1, 2, 3)")
    ap.add_argument("--lap", type=int, default=20, help="số lần lặp mỗi ca để đo T trong run_tests.py")
    ap.add_argument("--giai-doan", default="Thực nghiệm – kiểm thử tự động A/B")
    ap.add_argument("--muc-tieu", action="append", help="ghi mục tiêu riêng (dùng nhiều lần để có nhiều dòng)")
    ap.add_argument("--ghi-chu", default="", help="ghi chú ngắn của học sinh cho lần chạy này")
    ap.add_argument("--seed", type=int, default=2026, help="hạt giống khi phải sinh bộ dữ liệu mô phỏng")
    ap.add_argument("--sinh-lai", action="store_true", help="sinh lại bộ dữ liệu mô phỏng trước khi chạy")
    ap.add_argument("--kiem-tra", action="store_true", help="chỉ kiểm tra tính toàn vẹn của sổ rồi thoát")
    a = ap.parse_args()

    os.makedirs(os.path.join(THU_MUC_NK, "trang"), exist_ok=True)
    tt = doc_tt()
    van_de = kiem_tra_chuoi(tt)
    if a.kiem_tra:
        print(f"Sổ có {tt['so_trang']} trang.")
        print("\n".join(van_de) if van_de else "Chuỗi mã băm nguyên vẹn: không có trang nào bị sửa hay xoá.")
        return

    bat_dau = datetime.now()
    tep_ca = os.path.join(GOC, "du_lieu", "kiem_thu", "bo_ca_mo_rong.json")
    if a.sinh_lai or not os.path.exists(tep_ca):
        import sinh_du_lieu_kiem_thu
        sinh_du_lieu_kiem_thu.sinh(a.seed)
        print(f"Đã sinh bộ dữ liệu mô phỏng (seed {a.seed}).")
    print(f"Chạy 24 ca chuẩn × {a.lan} lần thử ...")
    kq24 = chay_24_ca(a.lan, a.lap)
    print(f"Chạy bộ ca mở rộng × {a.lan} lần thử ...")
    th_mr, dong_mr = chay_mo_rong(a.lan)
    bo_ca = json.load(open(tep_ca, encoding="utf-8"))

    p = lap_trang(a, tt, bat_dau, kq24, th_mr, dong_mr, bo_ca, van_de)
    ten = f"trang/{p['trang']:03d}_{bat_dau:%Y%m%d_%H%M%S}.json"
    noi_dung = json.dumps(p, ensure_ascii=False, indent=1).encode("utf-8")
    with open(os.path.join(THU_MUC_NK, ten), "xb") as f:  # "x": không bao giờ ghi đè trang đã có
        f.write(noi_dung)
    them_trang_md(p)
    them_trang_docx(p)
    tt["so_trang"] = p["trang"]
    tt["cac_trang"].append({"trang": p["trang"], "ngay": p["ngay"], "tep": ten, "sha256": sha(noi_dung),
                            "sha_trang_truoc": p["sha_trang_truoc"]})
    tt["sha_docx"] = sha_tep(TEP_DOCX)
    with open(TEP_TT, "w", encoding="utf-8") as f:
        json.dump(tt, f, ensure_ascii=False, indent=1)

    print(f"\nĐã ghi TRANG {p['trang']} ({p['ngay']}, {p['tu']}–{p['den']}) vào nhat_ky_nghien_cuu/")
    for h in p["bang24"]:
        print("  " + " | ".join(h))
    for x in p["hien_tuong"]:
        print("  * " + x)


if __name__ == "__main__":
    main()
